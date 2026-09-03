"""Unit tests for campaign artifact generation service (Phase 6)."""
from pathlib import Path

from openpyxl import load_workbook

from src.services.artifact_generation import ArtifactStatus, ArtifactType, CampaignArtifactService


def _service(tmp_path: Path) -> CampaignArtifactService:
    return CampaignArtifactService(base_dir=tmp_path / "artifacts")


def test_generates_all_required_artifacts(tmp_path: Path) -> None:
    service = _service(tmp_path)
    campaign_id = 101
    author = "planner.user"

    generators = [
        (service.generate_selection_briefing, ArtifactType.SELECTION_BRIEFING),
        (service.generate_audience_definition, ArtifactType.AUDIENCE_DEFINITION),
        (service.generate_campaign_configuration, ArtifactType.CAMPAIGN_CONFIGURATION),
        (service.generate_proposition_sheet, ArtifactType.PROPOSITION_SHEET),
        (service.generate_treatment_sheet, ArtifactType.TREATMENT_SHEET),
        (service.generate_contact_rules, ArtifactType.CONTACT_RULES),
        (service.generate_volume_constraints, ArtifactType.VOLUME_CONSTRAINTS),
        (service.generate_control_group_definition, ArtifactType.CONTROL_GROUP_DEFINITION),
        (service.generate_reporting_configuration, ArtifactType.REPORTING_CONFIGURATION),
    ]

    for generate, expected_type in generators:
        artifact = generate(campaign_id, {"key": expected_type.value}, author)
        assert artifact.version == 1
        assert artifact.artifact_type == expected_type
        assert artifact.status == ArtifactStatus.DRAFT


def test_versioning_approval_and_rollback(tmp_path: Path) -> None:
    service = _service(tmp_path)
    campaign_id = 202

    version1 = service.generate_selection_briefing(campaign_id, {"objective": "acquire"}, "alice")
    version2 = service.generate_selection_briefing(campaign_id, {"objective": "retain"}, "alice")

    assert version1.version == 1
    assert version2.version == 2

    pending = service.submit_for_approval(campaign_id, ArtifactType.SELECTION_BRIEFING)
    assert pending.status == ArtifactStatus.PENDING_APPROVAL

    approved = service.review_approval(
        campaign_id,
        ArtifactType.SELECTION_BRIEFING,
        reviewer="manager.bob",
        approve=True,
        comment="looks good",
    )
    assert approved.status == ArtifactStatus.APPROVED
    assert approved.approver == "manager.bob"

    rollback = service.rollback_to_version(
        campaign_id,
        ArtifactType.SELECTION_BRIEFING,
        target_version=1,
        actor="ops.carol",
    )
    assert rollback.version == 3
    assert rollback.status == ArtifactStatus.ROLLED_BACK
    assert rollback.content["objective"] == "acquire"


def test_excel_and_word_exports(tmp_path: Path) -> None:
    service = _service(tmp_path)
    campaign_id = 303

    service.generate_selection_briefing(campaign_id, {"segment": "A"}, "alice")
    service.generate_audience_definition(campaign_id, {"audience": "high_value"}, "alice")

    excel_path = service.export_excel(campaign_id)
    word_path = service.export_word(campaign_id)

    assert excel_path.exists()
    assert word_path.exists()

    workbook = load_workbook(excel_path)
    assert "Summary" in workbook.sheetnames
    assert "selection_briefing" in workbook.sheetnames

    assert word_path.suffix == ".docx"
