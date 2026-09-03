"""Campaign artifact generation with versioning, approvals, rollback, and exports."""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from uuid import uuid4

from docx import Document
from openpyxl import Workbook


class ArtifactType(str, Enum):
    SELECTION_BRIEFING = "selection_briefing"
    AUDIENCE_DEFINITION = "audience_definition"
    CAMPAIGN_CONFIGURATION = "campaign_configuration"
    PROPOSITION_SHEET = "proposition_sheet"
    TREATMENT_SHEET = "treatment_sheet"
    CONTACT_RULES = "contact_rules"
    VOLUME_CONSTRAINTS = "volume_constraints"
    CONTROL_GROUP_DEFINITION = "control_group_definition"
    REPORTING_CONFIGURATION = "reporting_configuration"


class ArtifactStatus(str, Enum):
    DRAFT = "draft"
    PENDING_APPROVAL = "pending_approval"
    APPROVED = "approved"
    REJECTED = "rejected"
    ROLLED_BACK = "rolled_back"


@dataclass
class ArtifactVersion:
    artifact_id: str
    campaign_id: int
    artifact_type: ArtifactType
    title: str
    version: int
    status: ArtifactStatus
    content: dict
    author: str
    created_at: str
    updated_at: str
    approver: str | None = None
    approval_comment: str | None = None

    def to_dict(self) -> dict:
        return {
            "artifact_id": self.artifact_id,
            "campaign_id": self.campaign_id,
            "artifact_type": self.artifact_type.value,
            "title": self.title,
            "version": self.version,
            "status": self.status.value,
            "content": self.content,
            "author": self.author,
            "created_at": self.created_at,
            "updated_at": self.updated_at,
            "approver": self.approver,
            "approval_comment": self.approval_comment,
        }

    @classmethod
    def from_dict(cls, payload: dict) -> "ArtifactVersion":
        return cls(
            artifact_id=payload["artifact_id"],
            campaign_id=payload["campaign_id"],
            artifact_type=ArtifactType(payload["artifact_type"]),
            title=payload["title"],
            version=payload["version"],
            status=ArtifactStatus(payload["status"]),
            content=payload.get("content", {}),
            author=payload["author"],
            created_at=payload["created_at"],
            updated_at=payload["updated_at"],
            approver=payload.get("approver"),
            approval_comment=payload.get("approval_comment"),
        )


@dataclass
class ArtifactHistory:
    current_version: int = 0
    versions: dict[int, ArtifactVersion] = field(default_factory=dict)


class CampaignArtifactService:
    """Produces and governs campaign artifacts through their lifecycle."""

    def __init__(self, base_dir: Path | None = None) -> None:
        root = Path(__file__).resolve().parents[4]
        self._base_dir = base_dir or root / "backend" / "logs" / "campaign_artifacts"
        self._base_dir.mkdir(parents=True, exist_ok=True)

    def _now(self) -> str:
        return datetime.now(timezone.utc).isoformat()

    def _campaign_dir(self, campaign_id: int) -> Path:
        path = self._base_dir / str(campaign_id)
        path.mkdir(parents=True, exist_ok=True)
        return path

    def _artifact_dir(self, campaign_id: int, artifact_type: ArtifactType) -> Path:
        path = self._campaign_dir(campaign_id) / artifact_type.value
        path.mkdir(parents=True, exist_ok=True)
        return path

    def _history_path(self, campaign_id: int, artifact_type: ArtifactType) -> Path:
        return self._artifact_dir(campaign_id, artifact_type) / "history.json"

    def _version_path(self, campaign_id: int, artifact_type: ArtifactType, version: int) -> Path:
        return self._artifact_dir(campaign_id, artifact_type) / f"v{version}.json"

    def _load_history(self, campaign_id: int, artifact_type: ArtifactType) -> ArtifactHistory:
        history_path = self._history_path(campaign_id, artifact_type)
        if not history_path.exists():
            return ArtifactHistory()

        payload = json.loads(history_path.read_text(encoding="utf-8"))
        versions = {
            int(key): ArtifactVersion.from_dict(value)
            for key, value in payload.get("versions", {}).items()
        }
        return ArtifactHistory(
            current_version=payload.get("current_version", 0),
            versions=versions,
        )

    def _save_history(self, campaign_id: int, artifact_type: ArtifactType, history: ArtifactHistory) -> None:
        payload = {
            "current_version": history.current_version,
            "versions": {str(key): value.to_dict() for key, value in history.versions.items()},
        }
        self._history_path(campaign_id, artifact_type).write_text(
            json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=True),
            encoding="utf-8",
        )

    def _save_version(self, artifact: ArtifactVersion) -> None:
        self._version_path(artifact.campaign_id, artifact.artifact_type, artifact.version).write_text(
            json.dumps(artifact.to_dict(), indent=2, sort_keys=True, ensure_ascii=True),
            encoding="utf-8",
        )

    def _next_version(self, history: ArtifactHistory) -> int:
        return history.current_version + 1

    def _create_or_update_artifact(
        self,
        campaign_id: int,
        artifact_type: ArtifactType,
        title: str,
        content: dict,
        author: str,
    ) -> ArtifactVersion:
        history = self._load_history(campaign_id, artifact_type)
        version = self._next_version(history)
        timestamp = self._now()

        artifact = ArtifactVersion(
            artifact_id=f"{artifact_type.value}-{uuid4().hex[:12]}",
            campaign_id=campaign_id,
            artifact_type=artifact_type,
            title=title,
            version=version,
            status=ArtifactStatus.DRAFT,
            content=content,
            author=author,
            created_at=timestamp,
            updated_at=timestamp,
        )
        history.versions[version] = artifact
        history.current_version = version
        self._save_version(artifact)
        self._save_history(campaign_id, artifact_type, history)
        return artifact

    def generate_artifact(
        self,
        campaign_id: int,
        artifact_type: ArtifactType,
        content: dict,
        author: str,
    ) -> ArtifactVersion:
        generators = {
            ArtifactType.SELECTION_BRIEFING: self.generate_selection_briefing,
            ArtifactType.AUDIENCE_DEFINITION: self.generate_audience_definition,
            ArtifactType.CAMPAIGN_CONFIGURATION: self.generate_campaign_configuration,
            ArtifactType.PROPOSITION_SHEET: self.generate_proposition_sheet,
            ArtifactType.TREATMENT_SHEET: self.generate_treatment_sheet,
            ArtifactType.CONTACT_RULES: self.generate_contact_rules,
            ArtifactType.VOLUME_CONSTRAINTS: self.generate_volume_constraints,
            ArtifactType.CONTROL_GROUP_DEFINITION: self.generate_control_group_definition,
            ArtifactType.REPORTING_CONFIGURATION: self.generate_reporting_configuration,
        }
        return generators[artifact_type](campaign_id, content, author)

    def generate_selection_briefing(self, campaign_id: int, content: dict, author: str) -> ArtifactVersion:
        return self._create_or_update_artifact(
            campaign_id,
            ArtifactType.SELECTION_BRIEFING,
            "Selection Briefing",
            content,
            author,
        )

    def generate_audience_definition(self, campaign_id: int, content: dict, author: str) -> ArtifactVersion:
        return self._create_or_update_artifact(
            campaign_id,
            ArtifactType.AUDIENCE_DEFINITION,
            "Audience Definition",
            content,
            author,
        )

    def generate_campaign_configuration(self, campaign_id: int, content: dict, author: str) -> ArtifactVersion:
        return self._create_or_update_artifact(
            campaign_id,
            ArtifactType.CAMPAIGN_CONFIGURATION,
            "Campaign Configuration",
            content,
            author,
        )

    def generate_proposition_sheet(self, campaign_id: int, content: dict, author: str) -> ArtifactVersion:
        return self._create_or_update_artifact(
            campaign_id,
            ArtifactType.PROPOSITION_SHEET,
            "Proposition Sheet",
            content,
            author,
        )

    def generate_treatment_sheet(self, campaign_id: int, content: dict, author: str) -> ArtifactVersion:
        return self._create_or_update_artifact(
            campaign_id,
            ArtifactType.TREATMENT_SHEET,
            "Treatment Sheet",
            content,
            author,
        )

    def generate_contact_rules(self, campaign_id: int, content: dict, author: str) -> ArtifactVersion:
        return self._create_or_update_artifact(
            campaign_id,
            ArtifactType.CONTACT_RULES,
            "Contact Rules",
            content,
            author,
        )

    def generate_volume_constraints(self, campaign_id: int, content: dict, author: str) -> ArtifactVersion:
        return self._create_or_update_artifact(
            campaign_id,
            ArtifactType.VOLUME_CONSTRAINTS,
            "Volume Constraints",
            content,
            author,
        )

    def generate_control_group_definition(self, campaign_id: int, content: dict, author: str) -> ArtifactVersion:
        return self._create_or_update_artifact(
            campaign_id,
            ArtifactType.CONTROL_GROUP_DEFINITION,
            "Control Group Definition",
            content,
            author,
        )

    def generate_reporting_configuration(self, campaign_id: int, content: dict, author: str) -> ArtifactVersion:
        return self._create_or_update_artifact(
            campaign_id,
            ArtifactType.REPORTING_CONFIGURATION,
            "Reporting Configuration",
            content,
            author,
        )

    def list_versions(self, campaign_id: int, artifact_type: ArtifactType) -> list[ArtifactVersion]:
        history = self._load_history(campaign_id, artifact_type)
        return [history.versions[key] for key in sorted(history.versions.keys())]

    def get_current_version(self, campaign_id: int, artifact_type: ArtifactType) -> ArtifactVersion | None:
        history = self._load_history(campaign_id, artifact_type)
        if history.current_version == 0:
            return None
        return history.versions.get(history.current_version)

    def submit_for_approval(self, campaign_id: int, artifact_type: ArtifactType) -> ArtifactVersion:
        history = self._load_history(campaign_id, artifact_type)
        if history.current_version == 0:
            raise ValueError("No artifact version exists for approval.")

        artifact = history.versions[history.current_version]
        artifact.status = ArtifactStatus.PENDING_APPROVAL
        artifact.updated_at = self._now()
        self._save_version(artifact)
        self._save_history(campaign_id, artifact_type, history)
        return artifact

    def submit_all_for_approval(self, campaign_id: int) -> list[ArtifactVersion]:
        submitted: list[ArtifactVersion] = []
        for artifact_type in ArtifactType:
            current = self.get_current_version(campaign_id, artifact_type)
            if current is None or current.status != ArtifactStatus.DRAFT:
                continue
            submitted.append(self.submit_for_approval(campaign_id, artifact_type))
        return submitted

    def review_approval(
        self,
        campaign_id: int,
        artifact_type: ArtifactType,
        reviewer: str,
        approve: bool,
        comment: str | None = None,
    ) -> ArtifactVersion:
        history = self._load_history(campaign_id, artifact_type)
        if history.current_version == 0:
            raise ValueError("No artifact version exists for review.")

        artifact = history.versions[history.current_version]
        if artifact.status != ArtifactStatus.PENDING_APPROVAL:
            raise ValueError("Artifact must be in pending_approval status before review.")

        artifact.status = ArtifactStatus.APPROVED if approve else ArtifactStatus.REJECTED
        artifact.approver = reviewer
        artifact.approval_comment = comment
        artifact.updated_at = self._now()
        self._save_version(artifact)
        self._save_history(campaign_id, artifact_type, history)
        return artifact

    def review_all_pending(
        self,
        campaign_id: int,
        reviewer: str,
        approve: bool,
        comment: str | None = None,
    ) -> list[ArtifactVersion]:
        reviewed: list[ArtifactVersion] = []
        for artifact_type in ArtifactType:
            current = self.get_current_version(campaign_id, artifact_type)
            if current is None or current.status != ArtifactStatus.PENDING_APPROVAL:
                continue
            reviewed.append(
                self.review_approval(
                    campaign_id=campaign_id,
                    artifact_type=artifact_type,
                    reviewer=reviewer,
                    approve=approve,
                    comment=comment,
                )
            )
        return reviewed

    def rollback_to_version(
        self,
        campaign_id: int,
        artifact_type: ArtifactType,
        target_version: int,
        actor: str,
    ) -> ArtifactVersion:
        history = self._load_history(campaign_id, artifact_type)
        if target_version not in history.versions:
            raise ValueError(f"Version {target_version} does not exist for {artifact_type.value}.")

        source = history.versions[target_version]
        new_version = self._next_version(history)
        timestamp = self._now()
        rolled_back = ArtifactVersion(
            artifact_id=f"{artifact_type.value}-{uuid4().hex[:12]}",
            campaign_id=campaign_id,
            artifact_type=artifact_type,
            title=source.title,
            version=new_version,
            status=ArtifactStatus.ROLLED_BACK,
            content=source.content,
            author=actor,
            created_at=timestamp,
            updated_at=timestamp,
            approver=source.approver,
            approval_comment=f"Rollback to v{target_version}",
        )

        history.versions[new_version] = rolled_back
        history.current_version = new_version
        self._save_version(rolled_back)
        self._save_history(campaign_id, artifact_type, history)
        return rolled_back

    def _all_current_artifacts(self, campaign_id: int) -> list[ArtifactVersion]:
        current: list[ArtifactVersion] = []
        for artifact_type in ArtifactType:
            artifact = self.get_current_version(campaign_id, artifact_type)
            if artifact is not None:
                current.append(artifact)
        return current

    def export_excel(self, campaign_id: int, output_path: Path | None = None) -> Path:
        artifacts = self._all_current_artifacts(campaign_id)
        export_dir = self._campaign_dir(campaign_id) / "exports"
        export_dir.mkdir(parents=True, exist_ok=True)
        destination = output_path or export_dir / f"campaign_{campaign_id}_artifacts.xlsx"

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Artifact Type", "Version", "Status", "Title", "Author", "Updated At"])

        for artifact in artifacts:
            summary.append(
                [
                    artifact.artifact_type.value,
                    artifact.version,
                    artifact.status.value,
                    artifact.title,
                    artifact.author,
                    artifact.updated_at,
                ]
            )

            sheet_name = artifact.artifact_type.value[:31]
            details = workbook.create_sheet(title=sheet_name)
            details.append(["Field", "Value"])
            details.append(["artifact_id", artifact.artifact_id])
            details.append(["title", artifact.title])
            details.append(["version", artifact.version])
            details.append(["status", artifact.status.value])
            details.append(["author", artifact.author])
            details.append(["approver", artifact.approver or ""])
            details.append(["approval_comment", artifact.approval_comment or ""])
            details.append(["content_json", json.dumps(artifact.content, ensure_ascii=True, sort_keys=True)])

        workbook.save(destination)
        return destination

    def export_word(self, campaign_id: int, output_path: Path | None = None) -> Path:
        artifacts = self._all_current_artifacts(campaign_id)
        export_dir = self._campaign_dir(campaign_id) / "exports"
        export_dir.mkdir(parents=True, exist_ok=True)
        destination = output_path or export_dir / f"campaign_{campaign_id}_artifacts.docx"

        document = Document()
        document.add_heading(f"Campaign {campaign_id} Artifact Pack", level=0)
        document.add_paragraph(f"Generated at: {self._now()}")

        for artifact in artifacts:
            document.add_heading(artifact.title, level=1)
            document.add_paragraph(f"Type: {artifact.artifact_type.value}")
            document.add_paragraph(f"Version: {artifact.version}")
            document.add_paragraph(f"Status: {artifact.status.value}")
            document.add_paragraph(f"Author: {artifact.author}")
            if artifact.approver:
                document.add_paragraph(f"Approver: {artifact.approver}")
            if artifact.approval_comment:
                document.add_paragraph(f"Approval Comment: {artifact.approval_comment}")
            document.add_paragraph("Content")
            document.add_paragraph(json.dumps(artifact.content, indent=2, sort_keys=True, ensure_ascii=True))

        document.save(destination)
        return destination
