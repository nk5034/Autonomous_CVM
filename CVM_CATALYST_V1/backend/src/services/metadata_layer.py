"""Enterprise metadata services for Phase 5."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


@dataclass(slots=True)
class _Column:
    name: str
    data_type: str | None = None
    description: str | None = None
    nullable: bool | None = None
    is_primary_key: bool = False
    is_foreign_key: bool = False
    references_table: str | None = None
    source_sheet: str | None = None


@dataclass(slots=True)
class _Table:
    name: str
    description: str | None = None
    columns: dict[str, _Column] = field(default_factory=dict)


@dataclass(slots=True)
class _ParsedMetadata:
    source: str
    file_path: str
    tables: dict[str, _Table]


class EnterpriseMetadataService:
    """Parses metadata workbooks and generates graph/SQL/report artifacts."""

    def __init__(self) -> None:
        repo_root = Path(__file__).resolve().parents[3]
        self._default_data_dictionary_path = repo_root / "knowledge" / "Data Dictionary.xlsx"
        self._default_database_model_path = repo_root / "knowledge" / "Database Model.xlsx"

    def parse_data_dictionary(self, file_path: str | None = None) -> dict[str, Any]:
        """Parse Data Dictionary metadata workbook into canonical table/column structures."""
        source_path = Path(file_path) if file_path else self._default_data_dictionary_path
        workbook = self._load_workbook(source_path)

        preferred_sheet = "Rich_Data_Dictionary"
        if preferred_sheet in workbook.sheetnames:
            sheet = workbook[preferred_sheet]
        else:
            sheet = workbook[workbook.sheetnames[0]]

        rows = self._sheet_rows(sheet)
        parsed = self.parse_data_dictionary_rows(rows, str(source_path), sheet.title)
        return self._to_response(parsed)

    def parse_database_model(self, file_path: str | None = None) -> dict[str, Any]:
        """Parse Database Model workbook into canonical table/column structures."""
        source_path = Path(file_path) if file_path else self._default_database_model_path
        workbook = self._load_workbook(source_path)

        tables: dict[str, _Table] = {}
        for sheet in workbook.worksheets:
            if sheet.title.strip().lower() in {"readme", "mapping_notes"}:
                continue
            rows = self._sheet_rows(sheet)
            parsed_part = self.parse_database_model_rows(rows, sheet.title)
            for table_name, table in parsed_part.items():
                existing = tables.setdefault(table_name, _Table(name=table_name))
                if table.description and not existing.description:
                    existing.description = table.description
                existing.columns.update(table.columns)

        parsed = _ParsedMetadata(
            source="database_model",
            file_path=str(source_path),
            tables=tables,
        )
        return self._to_response(parsed)

    def parse_data_dictionary_rows(
        self,
        rows: list[dict[str, Any]],
        file_path: str = "inline",
        source_sheet: str = "inline",
    ) -> _ParsedMetadata:
        """Parse data dictionary rows. Exposed for deterministic unit tests."""
        tables: dict[str, _Table] = {}
        for row in rows:
            normalized = {self._normalize_key(k): v for k, v in row.items()}
            table_name = self._coalesce(
                normalized,
                ["table_name", "table", "entity", "dataset", "source_table", "target_table"],
            )
            if not table_name:
                continue

            normalized_table = self._normalize_table_name(str(table_name))
            table = tables.setdefault(normalized_table, _Table(name=normalized_table))
            table.description = table.description or self._coalesce(
                normalized,
                ["table_description", "entity_description", "dataset_description"],
            )

            column_name = self._coalesce(
                normalized,
                [
                    "column_name",
                    "column",
                    "field_name",
                    "attribute",
                    "column_logical_name",
                ],
            )
            if not column_name:
                continue

            normalized_col = self._normalize_column_name(str(column_name))
            table.columns[normalized_col] = _Column(
                name=normalized_col,
                data_type=self._coalesce(
                    normalized,
                    ["data_type", "datatype", "type", "column_type"],
                ),
                description=self._coalesce(
                    normalized,
                    ["description", "column_description", "definition", "business_definition"],
                ),
                nullable=self._to_bool(self._coalesce(normalized, ["nullable", "is_nullable", "null_allowed"])),
                source_sheet=source_sheet,
            )

        return _ParsedMetadata(source="data_dictionary", file_path=file_path, tables=tables)

    def parse_database_model_rows(self, rows: list[dict[str, Any]], sheet_name: str) -> dict[str, _Table]:
        """Parse one database model sheet into one or more canonical tables."""
        table_name_from_sheet = self._normalize_table_name(sheet_name.removesuffix("_Data"))
        tables: dict[str, _Table] = {}

        metadata_keys = {
            "column_name",
            "column",
            "field_name",
            "attribute",
            "physical_column_name",
        }
        is_rowwise_metadata = any(
            any(self._normalize_key(str(k)) in metadata_keys for k in row.keys())
            for row in rows
        )

        if rows and not is_rowwise_metadata:
            table = tables.setdefault(table_name_from_sheet, _Table(name=table_name_from_sheet))
            ordered_headers: list[str] = []
            seen_headers: set[str] = set()
            for row in rows:
                for raw_header in row.keys():
                    header = self._normalize_column_name(str(raw_header))
                    if not header or header in seen_headers:
                        continue
                    seen_headers.add(header)
                    ordered_headers.append(header)

            for header in ordered_headers:
                sample_values = [row.get(header) for row in rows if row.get(header) is not None]
                normalized_col = self._normalize_column_name(header)
                is_pk = normalized_col in {"id", f"{table_name_from_sheet}_id"}
                is_fk = not is_pk and normalized_col.endswith("_id")
                ref_table = None
                if is_fk:
                    ref_table = self._normalize_table_name(normalized_col[:-3])

                table.columns[normalized_col] = _Column(
                    name=normalized_col,
                    data_type=self._infer_data_type(sample_values),
                    description=None,
                    nullable=None,
                    is_primary_key=is_pk,
                    is_foreign_key=is_fk,
                    references_table=ref_table,
                    source_sheet=sheet_name,
                )

            return tables

        for row in rows:
            normalized = {self._normalize_key(k): v for k, v in row.items()}
            table_name = self._coalesce(normalized, ["table_name", "table", "entity", "dataset"])
            normalized_table = self._normalize_table_name(str(table_name)) if table_name else table_name_from_sheet
            table = tables.setdefault(normalized_table, _Table(name=normalized_table))
            table.description = table.description or self._coalesce(
                normalized,
                ["table_description", "entity_description", "dataset_description"],
            )

            column_name = self._coalesce(
                normalized,
                ["column_name", "column", "field_name", "attribute", "physical_column_name"],
            )
            if not column_name:
                continue
            normalized_col = self._normalize_column_name(str(column_name))

            key_text = (self._coalesce(normalized, ["key_type", "key", "constraint"]) or "").lower()
            is_pk = "pk" in key_text or "primary" in key_text
            is_fk = "fk" in key_text or "foreign" in key_text

            fk_ref = self._coalesce(
                normalized,
                ["references_table", "reference_table", "parent_table", "fk_table"],
            )

            table.columns[normalized_col] = _Column(
                name=normalized_col,
                data_type=self._coalesce(
                    normalized,
                    ["data_type", "datatype", "type", "column_type"],
                ),
                description=self._coalesce(
                    normalized,
                    ["description", "column_description", "definition", "business_definition"],
                ),
                nullable=self._to_bool(self._coalesce(normalized, ["nullable", "is_nullable", "null_allowed"])),
                is_primary_key=is_pk,
                is_foreign_key=is_fk,
                references_table=self._normalize_table_name(str(fk_ref)) if fk_ref else None,
                source_sheet=sheet_name,
            )

        return tables

    def build_semantic_relationship_graph(
        self,
        data_dictionary: dict[str, Any],
        database_model: dict[str, Any],
    ) -> dict[str, Any]:
        """Build semantic graph linking business metadata to physical metadata."""
        dd_columns = self._flatten_columns(data_dictionary)
        db_columns = self._flatten_columns(database_model)

        nodes: dict[str, dict[str, Any]] = {}
        edges: list[dict[str, Any]] = []

        for col in dd_columns:
            node_id = f"dd:{col['table']}.{col['column']}"
            nodes[node_id] = {
                "id": node_id,
                "label": f"{col['table']}.{col['column']}",
                "node_type": "business_column",
                "table": col["table"],
                "column": col["column"],
            }

        for col in db_columns:
            node_id = f"db:{col['table']}.{col['column']}"
            nodes[node_id] = {
                "id": node_id,
                "label": f"{col['table']}.{col['column']}",
                "node_type": "physical_column",
                "table": col["table"],
                "column": col["column"],
            }

        db_lookup = {(c["table"], c["column"]): c for c in db_columns}
        db_by_column: dict[str, list[dict[str, Any]]] = {}
        for c in db_columns:
            db_by_column.setdefault(c["column"], []).append(c)

        for dd_col in dd_columns:
            exact = db_lookup.get((dd_col["table"], dd_col["column"]))
            if exact:
                edges.append(
                    {
                        "source": f"dd:{dd_col['table']}.{dd_col['column']}",
                        "target": f"db:{exact['table']}.{exact['column']}",
                        "relationship_type": "maps_to_exact",
                        "confidence": 1.0,
                        "metadata": {},
                    }
                )
                continue

            for similar in db_by_column.get(dd_col["column"], []):
                edges.append(
                    {
                        "source": f"dd:{dd_col['table']}.{dd_col['column']}",
                        "target": f"db:{similar['table']}.{similar['column']}",
                        "relationship_type": "maps_to_name_match",
                        "confidence": 0.75,
                        "metadata": {},
                    }
                )

        return self._graph_response("semantic_relationship", nodes, edges)

    def build_column_lineage_graph(
        self,
        data_dictionary: dict[str, Any],
        database_model: dict[str, Any],
    ) -> dict[str, Any]:
        """Build column-level lineage graph from business dictionary to physical model."""
        semantic = self.build_semantic_relationship_graph(data_dictionary, database_model)
        lineage_edges: list[dict[str, Any]] = []

        for edge in semantic["edges"]:
            lineage_edges.append(
                {
                    "source": edge["source"],
                    "target": edge["target"],
                    "relationship_type": "column_lineage",
                    "confidence": edge["confidence"],
                    "metadata": {"derived_from": edge["relationship_type"]},
                }
            )

        nodes = {node["id"]: node for node in semantic["nodes"]}
        return self._graph_response("column_lineage", nodes, lineage_edges)

    def build_table_relationship_graph(self, database_model: dict[str, Any]) -> dict[str, Any]:
        """Build table relationship graph using FK metadata and deterministic inference."""
        nodes: dict[str, dict[str, Any]] = {}
        edges: list[dict[str, Any]] = []

        tables = database_model.get("tables", [])
        table_names = {t["name"] for t in tables}

        for table in tables:
            node_id = f"table:{table['name']}"
            nodes[node_id] = {
                "id": node_id,
                "label": table["name"],
                "node_type": "table",
                "table": table["name"],
                "column": None,
            }

        for table in tables:
            source = table["name"]
            for col in table.get("columns", []):
                ref = col.get("references_table")
                if col.get("is_foreign_key") and ref:
                    if ref in table_names and source != ref:
                        edges.append(
                            {
                                "source": f"table:{source}",
                                "target": f"table:{ref}",
                                "relationship_type": "foreign_key",
                                "confidence": 1.0,
                                "metadata": {"fk_column": col["name"]},
                            }
                        )
                    continue

                inferred = self._infer_reference_table(col.get("name", ""), table_names)
                if inferred and inferred != source:
                    edges.append(
                        {
                            "source": f"table:{source}",
                            "target": f"table:{inferred}",
                            "relationship_type": "inferred_foreign_key",
                            "confidence": 0.7,
                            "metadata": {"fk_column": col["name"]},
                        }
                    )

        return self._graph_response("table_relationship", nodes, self._unique_edges(edges))

    def generate_sql(
        self,
        base_table: str,
        select_columns: list[str],
        filters: dict[str, Any],
        table_relationship_graph: dict[str, Any],
        limit: int = 100,
    ) -> dict[str, Any]:
        """Generate SQL from selected columns and table relationship graph."""
        normalized_base = self._normalize_table_name(base_table)
        join_edges = table_relationship_graph.get("edges", [])

        selected = select_columns or [f"{normalized_base}.*"]
        target_tables = set()
        unresolved: list[str] = []

        for item in selected:
            if "." not in item:
                continue
            table, _ = item.split(".", 1)
            table = self._normalize_table_name(table)
            if table != normalized_base:
                target_tables.add(table)

        joins: list[str] = []
        joined_tables = [normalized_base]

        for table in sorted(target_tables):
            edge = self._find_join_edge(normalized_base, table, join_edges)
            if not edge:
                unresolved.append(f"{table}.*")
                continue
            join_sql = self._build_join_clause(normalized_base, table, edge)
            joins.append(join_sql)
            joined_tables.append(table)

        where_sql = ""
        if filters:
            conditions = [
                f"{k} = {self._sql_literal(v)}"
                for k, v in sorted(filters.items(), key=lambda item: item[0])
            ]
            where_sql = "\nWHERE " + " AND ".join(conditions)

        select_sql = ", ".join(selected)
        join_sql = "\n".join(joins)
        sql = (
            f"SELECT {select_sql}\n"
            f"FROM {normalized_base}\n"
            f"{join_sql}{where_sql}\n"
            f"LIMIT {max(1, int(limit))};"
        ).strip()

        return {
            "sql": sql,
            "joined_tables": joined_tables,
            "unresolved_columns": unresolved,
        }

    def generate_waterfall_analysis(
        self,
        base_table: str,
        entity_id_column: str,
        stages: list[dict[str, str]],
    ) -> dict[str, Any]:
        """Generate SQL for sequential waterfall (funnel) analysis."""
        normalized_base = self._normalize_table_name(base_table)
        normalized_id = self._normalize_column_name(entity_id_column)

        if not stages:
            sql = (
                f"SELECT COUNT(DISTINCT {normalized_id}) AS total_entities\n"
                f"FROM {normalized_base};"
            )
            return {"sql": sql, "stage_count": 0, "stage_names": []}

        ctes = [
            f"base AS (\n"
            f"    SELECT DISTINCT {normalized_id} AS entity_id\n"
            f"    FROM {normalized_base}\n"
            f")"
        ]

        union_parts = ["SELECT 'base' AS stage_name, COUNT(*) AS stage_count FROM base"]
        previous_stage = "base"

        for index, stage in enumerate(stages, start=1):
            stage_alias = f"stage_{index}"
            condition = stage["condition"].strip()
            ctes.append(
                f"{stage_alias} AS (\n"
                f"    SELECT b.entity_id\n"
                f"    FROM {previous_stage} b\n"
                f"    JOIN {normalized_base} t ON t.{normalized_id} = b.entity_id\n"
                f"    WHERE {condition}\n"
                f")"
            )
            union_parts.append(
                f"SELECT '{stage['name']}' AS stage_name, COUNT(*) AS stage_count FROM {stage_alias}"
            )
            previous_stage = stage_alias

        sql = "WITH\n" + ",\n".join(ctes) + "\n" + "\nUNION ALL\n".join(union_parts) + ";"
        return {
            "sql": sql,
            "stage_count": len(stages),
            "stage_names": [stage["name"] for stage in stages],
        }

    def generate_explainability_report(
        self,
        sql: str,
        semantic_graph: dict[str, Any],
        lineage_graph: dict[str, Any],
        table_relationship_graph: dict[str, Any],
    ) -> dict[str, Any]:
        """Generate explainability report for generated SQL and metadata graphs."""
        semantic_edges = semantic_graph.get("edge_count", len(semantic_graph.get("edges", [])))
        lineage_edges = lineage_graph.get("edge_count", len(lineage_graph.get("edges", [])))
        table_edges = table_relationship_graph.get("edge_count", len(table_relationship_graph.get("edges", [])))

        assumptions = [
            "Data Dictionary and Database Model workbooks are treated as authoritative metadata sources.",
            "Column name normalization uses lowercase and underscores for deterministic matching.",
            "Inferred foreign keys rely on <table>_id naming conventions when explicit FK metadata is absent.",
        ]

        summary = (
            "Generated SQL is traceable to enterprise metadata through semantic mappings, "
            "column lineage links, and table relationships."
        )

        return {
            "summary": summary,
            "assumptions": assumptions,
            "semantic_coverage": {
                "nodes": semantic_graph.get("node_count", len(semantic_graph.get("nodes", []))),
                "edges": semantic_edges,
            },
            "lineage_coverage": {
                "nodes": lineage_graph.get("node_count", len(lineage_graph.get("nodes", []))),
                "edges": lineage_edges,
            },
            "relationship_coverage": {
                "nodes": table_relationship_graph.get(
                    "node_count",
                    len(table_relationship_graph.get("nodes", [])),
                ),
                "edges": table_edges,
            },
            "sql_preview": sql.strip().splitlines()[:5],
        }

    def _to_response(self, parsed: _ParsedMetadata) -> dict[str, Any]:
        tables = []
        column_count = 0
        for table_name in sorted(parsed.tables):
            table = parsed.tables[table_name]
            columns = []
            for column_name in sorted(table.columns):
                col = table.columns[column_name]
                columns.append(
                    {
                        "name": col.name,
                        "data_type": col.data_type,
                        "description": col.description,
                        "nullable": col.nullable,
                        "is_primary_key": col.is_primary_key,
                        "is_foreign_key": col.is_foreign_key,
                        "references_table": col.references_table,
                        "source_sheet": col.source_sheet,
                    }
                )
            column_count += len(columns)
            tables.append(
                {
                    "name": table.name,
                    "description": table.description,
                    "columns": columns,
                }
            )

        return {
            "source": parsed.source,
            "file_path": parsed.file_path,
            "table_count": len(tables),
            "column_count": column_count,
            "tables": tables,
        }

    def _graph_response(
        self,
        graph_type: str,
        nodes: dict[str, dict[str, Any]],
        edges: list[dict[str, Any]],
    ) -> dict[str, Any]:
        node_values = [nodes[k] for k in sorted(nodes.keys())]
        edge_values = sorted(
            edges,
            key=lambda item: (item["source"], item["target"], item["relationship_type"]),
        )
        return {
            "graph_type": graph_type,
            "node_count": len(node_values),
            "edge_count": len(edge_values),
            "nodes": node_values,
            "edges": edge_values,
        }

    def _flatten_columns(self, metadata: dict[str, Any]) -> list[dict[str, Any]]:
        flattened: list[dict[str, Any]] = []
        for table in metadata.get("tables", []):
            table_name = self._normalize_table_name(table["name"])
            for col in table.get("columns", []):
                flattened.append(
                    {
                        "table": table_name,
                        "column": self._normalize_column_name(col["name"]),
                        "is_primary_key": col.get("is_primary_key", False),
                        "is_foreign_key": col.get("is_foreign_key", False),
                        "references_table": col.get("references_table"),
                    }
                )
        return flattened

    def _find_join_edge(
        self,
        base_table: str,
        target_table: str,
        edges: list[dict[str, Any]],
    ) -> dict[str, Any] | None:
        base_node = f"table:{base_table}"
        target_node = f"table:{target_table}"

        for edge in edges:
            if edge["source"] == base_node and edge["target"] == target_node:
                return edge
            if edge["source"] == target_node and edge["target"] == base_node:
                return edge
        return None

    def _infer_data_type(self, values: list[Any]) -> str | None:
        if not values:
            return None

        first = values[0]
        type_name = type(first).__name__.lower()

        if "datetime" in type_name:
            return "datetime"
        if isinstance(first, bool):
            return "boolean"
        if isinstance(first, int):
            return "integer"
        if isinstance(first, float):
            return "float"
        return "string"

    def _build_join_clause(self, base_table: str, target_table: str, edge: dict[str, Any]) -> str:
        fk_column = edge.get("metadata", {}).get("fk_column")
        source = edge.get("source", "")

        if fk_column:
            if source == f"table:{base_table}":
                condition = (
                    f"{base_table}.{self._normalize_column_name(fk_column)} = "
                    f"{target_table}.id"
                )
            else:
                condition = (
                    f"{target_table}.{self._normalize_column_name(fk_column)} = "
                    f"{base_table}.id"
                )
        else:
            condition = f"{target_table}.id = {base_table}.id"

        return f"LEFT JOIN {target_table} ON {condition}"

    def _unique_edges(self, edges: list[dict[str, Any]]) -> list[dict[str, Any]]:
        seen: set[tuple[str, str, str, str]] = set()
        deduped: list[dict[str, Any]] = []
        for edge in edges:
            fk = edge.get("metadata", {}).get("fk_column", "")
            key = (edge["source"], edge["target"], edge["relationship_type"], fk)
            if key in seen:
                continue
            seen.add(key)
            deduped.append(edge)
        return deduped

    def _infer_reference_table(self, column_name: str, table_names: set[str]) -> str | None:
        normalized = self._normalize_column_name(column_name)
        if not normalized.endswith("_id"):
            return None

        stem = normalized[:-3]
        candidates = [stem, f"{stem}s"]
        for candidate in candidates:
            normalized_candidate = self._normalize_table_name(candidate)
            if normalized_candidate in table_names:
                return normalized_candidate
        return None

    def _sheet_rows(self, sheet: Any) -> list[dict[str, Any]]:
        values = list(sheet.iter_rows(values_only=True))
        non_empty_rows = [row for row in values if any(cell is not None and str(cell).strip() for cell in row)]
        if not non_empty_rows:
            return []

        header_idx = self._find_header_row_index(non_empty_rows)
        headers = [
            self._normalize_key(str(cell)) if cell is not None else ""
            for cell in non_empty_rows[header_idx]
        ]
        rows: list[dict[str, Any]] = []

        for row in non_empty_rows[header_idx + 1 :]:
            item: dict[str, Any] = {}
            for idx, value in enumerate(row):
                header = headers[idx] if idx < len(headers) else f"column_{idx + 1}"
                if not header:
                    continue
                item[header] = value
            if item:
                rows.append(item)

        return rows

    def _find_header_row_index(self, rows: list[tuple[Any, ...]]) -> int:
        best_idx = 0
        best_score = -1

        for idx, row in enumerate(rows):
            score = 0
            for cell in row:
                if not isinstance(cell, str):
                    continue
                text = cell.strip()
                if not text:
                    continue
                normalized = self._normalize_key(text)
                if not normalized:
                    continue
                if all(ch.isalnum() or ch == "_" for ch in normalized):
                    score += 1

            if score > best_score:
                best_idx = idx
                best_score = score

        return best_idx

    def _load_workbook(self, path: Path) -> Any:
        if not path.exists():
            raise FileNotFoundError(f"Metadata workbook not found: {path}")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl is required to parse .xlsx metadata files. Install it in backend dependencies."
            ) from exc
        return load_workbook(path, read_only=True, data_only=True)

    def _normalize_key(self, key: str) -> str:
        return "_".join(str(key).strip().lower().replace("-", " ").split())

    def _normalize_table_name(self, name: str) -> str:
        return self._normalize_key(name)

    def _normalize_column_name(self, name: str) -> str:
        return self._normalize_key(name)

    def _coalesce(self, row: dict[str, Any], keys: list[str]) -> Any:
        for key in keys:
            if key in row:
                value = row[key]
                if value is None:
                    continue
                if isinstance(value, str) and not value.strip():
                    continue
                return value
        return None

    def _to_bool(self, value: Any) -> bool | None:
        if value is None:
            return None
        if isinstance(value, bool):
            return value
        text = str(value).strip().lower()
        if text in {"true", "t", "yes", "y", "1"}:
            return True
        if text in {"false", "f", "no", "n", "0"}:
            return False
        return None

    def _sql_literal(self, value: Any) -> str:
        if value is None:
            return "NULL"
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, (int, float)):
            return str(value)
        escaped = str(value).replace("'", "''")
        return f"'{escaped}'"
