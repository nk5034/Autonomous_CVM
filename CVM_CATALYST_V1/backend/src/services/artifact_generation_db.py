"""Database-backed campaign artifact lifecycle service."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from src.models.entities import (
    CampaignArtifact,
    CampaignArtifactApproval,
    CampaignArtifactStatus,
    CampaignArtifactType,
)
from src.repositories.campaign_artifacts import (
    SqlAlchemyCampaignArtifactApprovalRepository,
    SqlAlchemyCampaignArtifactRepository,
)
from src.repositories.interfaces import (
    CampaignArtifactApprovalRepository,
    CampaignArtifactRepository,
)


class CampaignArtifactDbService:
    """Async artifact generation service backed by SQLAlchemy repositories."""

    def __init__(
        self,
        session: AsyncSession,
        base_dir: Path | None = None,
        artifact_repo: CampaignArtifactRepository | None = None,
        approval_repo: CampaignArtifactApprovalRepository | None = None,
    ) -> None:
        self._session = session
        self._artifact_repo = artifact_repo or SqlAlchemyCampaignArtifactRepository(session)
        self._approval_repo = approval_repo or SqlAlchemyCampaignArtifactApprovalRepository(session)
        root = Path(__file__).resolve().parents[4]
        self._export_dir = base_dir or root / "backend" / "logs" / "campaign_artifacts_db"
        self._export_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _title_for(artifact_type: CampaignArtifactType) -> str:
        titles = {
            CampaignArtifactType.SELECTION_BRIEFING: "Selection Briefing",
            CampaignArtifactType.AUDIENCE_DEFINITION: "Audience Definition",
            CampaignArtifactType.CAMPAIGN_CONFIGURATION: "Campaign Configuration",
            CampaignArtifactType.PROPOSITION_SHEET: "Proposition Sheet",
            CampaignArtifactType.TREATMENT_SHEET: "Treatment Sheet",
            CampaignArtifactType.CONTACT_RULES: "Contact Rules",
            CampaignArtifactType.VOLUME_CONSTRAINTS: "Volume Constraints",
            CampaignArtifactType.CONTROL_GROUP_DEFINITION: "Control Group Definition",
            CampaignArtifactType.REPORTING_CONFIGURATION: "Reporting Configuration",
        }
        return titles[artifact_type]

    async def generate_artifact(
        self,
        campaign_id: int,
        artifact_type: CampaignArtifactType,
        content: dict,
        author: str,
    ) -> CampaignArtifact:
        version = await self._artifact_repo.next_version(campaign_id, artifact_type)
        artifact = CampaignArtifact(
            campaign_id=campaign_id,
            artifact_type=artifact_type,
            title=self._title_for(artifact_type),
            version=version,
            status=CampaignArtifactStatus.DRAFT,
            content=content,
            author=author,
        )
        await self._artifact_repo.create(artifact)
        await self._session.commit()
        await self._session.refresh(artifact)
        return artifact

    async def list_versions(self, campaign_id: int, artifact_type: CampaignArtifactType) -> list[CampaignArtifact]:
        return await self._artifact_repo.list_versions(campaign_id, artifact_type)

    async def get_current(self, campaign_id: int, artifact_type: CampaignArtifactType) -> CampaignArtifact | None:
        return await self._artifact_repo.get_current(campaign_id, artifact_type)

    async def submit_for_approval(self, campaign_id: int, artifact_type: CampaignArtifactType) -> CampaignArtifact:
        artifact = await self._artifact_repo.get_current(campaign_id, artifact_type)
        if artifact is None:
            raise ValueError("No artifact version exists for approval.")
        artifact.status = CampaignArtifactStatus.PENDING_APPROVAL
        await self._artifact_repo.save(artifact)
        await self._session.commit()
        await self._session.refresh(artifact)
        return artifact

    async def review_approval(
        self,
        campaign_id: int,
        artifact_type: CampaignArtifactType,
        reviewer: str,
        approve: bool,
        comment: str | None = None,
    ) -> CampaignArtifact:
        artifact = await self._artifact_repo.get_current(campaign_id, artifact_type)
        if artifact is None:
            raise ValueError("No artifact version exists for review.")
        if artifact.status != CampaignArtifactStatus.PENDING_APPROVAL:
            raise ValueError("Artifact must be in pending_approval status before review.")

        artifact.status = CampaignArtifactStatus.APPROVED if approve else CampaignArtifactStatus.REJECTED
        artifact.approver = reviewer
        artifact.approval_comment = comment

        await self._approval_repo.add(
            CampaignArtifactApproval(
                artifact_id=artifact.id,
                reviewer=reviewer,
                approved=approve,
                comment=comment,
            )
        )
        await self._artifact_repo.save(artifact)
        await self._session.commit()
        await self._session.refresh(artifact)
        return artifact

    async def rollback_to_version(
        self,
        campaign_id: int,
        artifact_type: CampaignArtifactType,
        target_version: int,
        actor: str,
    ) -> CampaignArtifact:
        versions = await self._artifact_repo.list_versions(campaign_id, artifact_type)
        source = next((item for item in versions if item.version == target_version), None)
        if source is None:
            raise ValueError(f"Version {target_version} does not exist for {artifact_type.value}.")

        next_version = await self._artifact_repo.next_version(campaign_id, artifact_type)
        artifact = CampaignArtifact(
            campaign_id=campaign_id,
            artifact_type=artifact_type,
            title=source.title,
            version=next_version,
            status=CampaignArtifactStatus.ROLLED_BACK,
            content=source.content,
            author=actor,
            approver=source.approver,
            approval_comment=f"Rollback to v{target_version}",
        )
        await self._artifact_repo.create(artifact)
        await self._session.commit()
        await self._session.refresh(artifact)
        return artifact

    async def _current_artifacts(self, campaign_id: int) -> list[CampaignArtifact]:
        artifacts: list[CampaignArtifact] = []
        for artifact_type in CampaignArtifactType:
            current = await self._artifact_repo.get_current(campaign_id, artifact_type)
            if current is not None:
                artifacts.append(current)
        return artifacts

    async def export_excel(self, campaign_id: int) -> Path:
        artifacts = await self._current_artifacts(campaign_id)
        export_dir = self._export_dir / str(campaign_id)
        export_dir.mkdir(parents=True, exist_ok=True)
        destination = export_dir / f"campaign_{campaign_id}_artifacts.xlsx"

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Artifact Type", "Version", "Status", "Title", "Author", "Updated At"])

        for artifact in artifacts:
            summary.append(
                [
                    artifact.artifact_type.value,
                    artifact.version,
                    artifact.status.value,
                    artifact.title,
                    artifact.author,
                    artifact.updated_at.isoformat(),
                ]
            )
            details = workbook.create_sheet(title=artifact.artifact_type.value[:31])
            details.append(["Field", "Value"])
            details.append(["id", artifact.id])
            details.append(["version", artifact.version])
            details.append(["status", artifact.status.value])
            details.append(["content_json", json.dumps(artifact.content, ensure_ascii=True, sort_keys=True)])

        workbook.save(destination)
        return destination

    async def export_word(self, campaign_id: int) -> Path:
        artifacts = await self._current_artifacts(campaign_id)
        export_dir = self._export_dir / str(campaign_id)
        export_dir.mkdir(parents=True, exist_ok=True)
        destination = export_dir / f"campaign_{campaign_id}_artifacts.docx"

        document = Document()
        document.add_heading(f"Campaign {campaign_id} Artifact Pack", level=0)

        for artifact in artifacts:
            document.add_heading(artifact.title, level=1)
            document.add_paragraph(f"Type: {artifact.artifact_type.value}")
            document.add_paragraph(f"Version: {artifact.version}")
            document.add_paragraph(f"Status: {artifact.status.value}")
            document.add_paragraph(json.dumps(artifact.content, indent=2, sort_keys=True, ensure_ascii=True))

        document.save(destination)
        return destination
